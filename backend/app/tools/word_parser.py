"""
Word文件解析器
"""

from typing import List, Dict
from .base_tool import BaseParserTool, DocumentChunk, DocumentElement
import os
import logging

logger = logging.getLogger(__name__)


class WordParserTool(BaseParserTool):
    """Word解析器 - 支持.docx文件"""

    def __init__(self, llm_client=None, **kwargs):
        self.llm = llm_client
        from .table_processor import TableProcessor
        self.processor = TableProcessor(**kwargs)

    @property
    def supported_extensions(self) -> List[str]:
        return [".docx"]

    def validate(self, file_path: str) -> bool:
        return file_path.lower().endswith(".docx") and os.path.exists(file_path)

    async def parse(self, file_path: str, original_filename: str = None) -> List[DocumentChunk]:
        elements = self.parse_with_structure(file_path)
        return await self._elements_to_chunks(elements, file_path, original_filename)

    @staticmethod
    def _strip_uuid_prefix(filename: str) -> str:
        basename = os.path.basename(filename)
        if "_" in basename:
            prefix = basename.split("_", 1)[0]
            if len(prefix) == 8 and all(c in "0123456789abcdef" for c in prefix):
                return basename.split("_", 1)[1]
        return basename

    def _extract_text_from_image(self, image_data: bytes) -> str:
        import tempfile
        with tempfile.NamedTemporaryFile(delete=False, suffix=".png") as f:
            f.write(image_data)
            temp_path = f.name
        try:
            try:
                from paddleocr import PaddleOCR
                ocr = PaddleOCR(use_angle_cls=True, lang="ch")
                result = ocr.ocr(temp_path, cls=True)
                if result and result[0]:
                    text = "\n".join([line[1][0] for line in result[0]])
                    if text.strip():
                        return text
            except (ImportError, Exception):
                pass
            try:
                import pytesseract
                from PIL import Image
                image = Image.open(temp_path)
                text = pytesseract.image_to_string(image, lang='chi_sim+eng')
                if text.strip():
                    return text
            except (ImportError, Exception):
                pass
        finally:
            if os.path.exists(temp_path):
                os.remove(temp_path)
        return ""

    async def _describe_image_with_llm(self, image_data: bytes, context: str = "") -> str:
        if not self.llm:
            return ""
        try:
            prompt = "请详细描述这张图片的内容。如果图片中有人物，请描述其外貌特征；如果有文字，请提取所有文字内容；如果有设备或场景，请描述其特征。"
            if context:
                prompt += f"\n上下文信息：{context}"
            description = await self.llm.generate_with_image(prompt, image_data)
            return description
        except Exception as e:
            logger.warning(f"LLM图片描述失败: {e}")
            return ""

    def parse_with_structure(self, file_path: str) -> List[DocumentElement]:
        from docx import Document
        from lxml import etree

        doc = Document(file_path)
        elements = []

        for para_idx, paragraph in enumerate(doc.paragraphs):
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = paragraph.style.name if paragraph.style else ""
            is_heading = "Heading" in style_name or "标题" in style_name
            lines = text.split('\n')
            for line_idx, line in enumerate(lines):
                line = line.strip()
                if line:
                    elements.append(DocumentElement(
                        type="heading" if is_heading else "paragraph",
                        text=line, paragraph_index=para_idx, line_number=line_idx + 1,
                        metadata={"style": style_name, "is_heading": is_heading}
                    ))

        nsmap = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        txbx_contents = doc.element.body.findall('.//w:txbxContent', nsmap)
        seen_texts = set()
        for txbx in txbx_contents:
            paras = txbx.findall('.//w:p', nsmap)
            for p in paras:
                runs = p.findall('.//w:r', nsmap)
                text_parts = []
                for r in runs:
                    for t in r.findall('.//w:t', nsmap):
                        if t.text:
                            text_parts.append(t.text)
                text = ''.join(text_parts).strip()
                if text and text not in seen_texts:
                    seen_texts.add(text)
                    elements.append(DocumentElement(
                        type="paragraph", text=text, metadata={"source": "textbox"}
                    ))

        for table_idx, table in enumerate(doc.tables):
            headers = []
            rows = []
            for row_idx, row in enumerate(table.rows):
                cells = [cell.text.strip() for cell in row.cells]
                if row_idx == 0:
                    headers = cells
                else:
                    rows.append(cells)
            if headers:
                elements.append(DocumentElement(
                    type="table", headers=headers, rows=rows,
                    metadata={"table_index": table_idx}
                ))

        for rel in doc.part.rels.values():
            if "image" in rel.reltype:
                elements.append(DocumentElement(
                    type="image", image_data=rel.target_part.blob,
                    metadata={"rel_type": rel.reltype}
                ))

        return elements

    async def _elements_to_chunks(
        self, elements: List[DocumentElement], source_file: str, original_filename: str = None
    ) -> List[DocumentChunk]:
        chunks = []
        current_text = ""
        current_metadata = {}
        chunk_index = 0
        current_section = None
        if original_filename:
            file_name = original_filename
        else:
            file_name = self._strip_uuid_prefix(source_file)
        file_type = "docx"

        for element in elements:
            if element.type == "heading":
                if current_text.strip():
                    chunks.append(DocumentChunk.create(
                        content=current_text.strip(), source_file=file_name, file_type=file_type,
                        chunk_index=chunk_index, paragraph_index=current_metadata.get("paragraph_index"),
                        section_title=current_section, content_type="text"
                    ))
                    chunk_index += 1
                    current_text = ""
                current_section = element.text
            elif element.type == "table":
                if current_text.strip():
                    chunks.append(DocumentChunk.create(
                        content=current_text.strip(), source_file=file_name, file_type=file_type,
                        chunk_index=chunk_index, paragraph_index=current_metadata.get("paragraph_index"),
                        section_title=current_section, content_type="text"
                    ))
                    chunk_index += 1
                    current_text = ""
                table_chunks = self.processor.process_table(
                    headers=element.headers, rows=element.rows,
                    source_file=file_name, file_type=file_type,
                    chunk_index=chunk_index, section_title=current_section,
                )
                chunks.extend(table_chunks)
                chunk_index += len(table_chunks)
            elif element.type == "image":
                if current_text.strip():
                    chunks.append(DocumentChunk.create(
                        content=current_text.strip(), source_file=file_name, file_type=file_type,
                        chunk_index=chunk_index, paragraph_index=current_metadata.get("paragraph_index"),
                        section_title=current_section, content_type="text"
                    ))
                    chunk_index += 1
                    current_text = ""
                image_content = "[图片内容待提取]"
                if element.image_data:
                    try:
                        ocr_text = self._extract_text_from_image(element.image_data)
                        if ocr_text:
                            image_content = f"[图片OCR文字]\n{ocr_text}"
                            logger.info(f"Word图片OCR提取成功，{len(ocr_text)} 字符")
                        else:
                            logger.info("Word图片OCR未提取到文字，尝试LLM描述")
                            llm_desc = await self._describe_image_with_llm(element.image_data)
                            if llm_desc:
                                image_content = f"[图片内容]\n{llm_desc}"
                                logger.info(f"Word图片LLM描述成功，{len(llm_desc)} 字符")
                    except Exception as e:
                        logger.warning(f"Word图片处理失败: {e}")
                chunks.append(DocumentChunk.create(
                    content=image_content, source_file=file_name, file_type=file_type,
                    chunk_index=chunk_index, content_type="image", section_title=current_section
                ))
                chunk_index += 1
            else:
                current_text += element.text + "\n"
                if not current_metadata:
                    current_metadata = {"paragraph_index": element.paragraph_index}

        if current_text.strip():
            chunks.append(DocumentChunk.create(
                content=current_text.strip(), source_file=file_name, file_type=file_type,
                chunk_index=chunk_index, paragraph_index=current_metadata.get("paragraph_index"),
                section_title=current_section, content_type="text"
            ))
        return chunks
