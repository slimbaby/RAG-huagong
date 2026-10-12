"""
制造业工具集 - 订单查询、库存查询、质检查询

只读工具，数据来自 SQLite 数据库。
优先从 Excel 文件加载数据（backend/app/data/*.xlsx），
Excel 不存在时回退到内存模拟数据。
"""

import json
import logging
import os
import sqlite3
from datetime import datetime, timedelta
from pathlib import Path

logger = logging.getLogger(__name__)

_DB_PATH = "./checkpoints/manufacturing.db"
_db_ready = False

# Excel 数据目录 (相对于本文件)
_DATA_DIR = Path(__file__).resolve().parent.parent / "data"

# ── 回退用模拟数据配置 ──
_CUSTOMERS = ["浙江化工集团", "江苏新材料科技", "上海精细化工", "广东石化设备", "山东恒通化工"]
_PRODUCTS = ["聚乙烯颗粒", "聚丙烯颗粒", "PVC树脂", "环氧树脂", "丙烯酸丁酯"]
_MATERIALS = [
    ("M001", "乙烯", "吨", 500),
    ("M002", "丙烯", "吨", 300),
    ("M003", "氯乙烯单体", "吨", 200),
    ("M004", "环氧氯丙烷", "桶", 150),
    ("M005", "引发剂A", "箱", 80),
    ("M006", "阻聚剂B", "箱", 120),
    ("M007", "稳定剂C", "袋", 200),
    ("M008", "包装袋(25kg)", "个", 5000),
    ("M009", "托盘", "个", 300),
    ("M010", "标签纸", "卷", 1000),
]
_WAREHOUSES = ["原料仓A", "成品仓B", "辅料仓C"]
_INSPECTION_ITEMS = {
    "聚乙烯颗粒": ["熔融指数", "密度", "拉伸强度", "灰分"],
    "聚丙烯颗粒": ["熔融指数", "等规度", "拉伸强度", "灰分"],
    "PVC树脂": ["聚合度", "挥发分", "表观密度", "筛余物"],
    "环氧树脂": ["环氧值", "软化点", "色泽", "挥发分"],
    "丙烯酸丁酯": ["纯度", "酸值", "水分", "色泽"],
}
_INSPECTORS = ["张工", "李工", "王工", "赵工"]
_ORDER_STATUSES = ["待采购", "生产中", "已发货", "已完成", "已取消"]


def _ensure_db() -> bool:
    """首次调用时创建 SQLite 数据库，优先从 Excel 加载数据。"""
    global _db_ready
    if _db_ready:
        return True
    try:
        os.makedirs(os.path.dirname(_DB_PATH), exist_ok=True)
        if os.path.exists(_DB_PATH):
            os.remove(_DB_PATH)
        conn = sqlite3.connect(_DB_PATH)
        conn.execute("PRAGMA journal_mode=WAL")
        _create_tables(conn)

        # 优先从 Excel 加载
        if _load_from_excel(conn):
            logger.info(f"[manufacturing_tools] 从 Excel 加载数据完成: {_DB_PATH}")
        else:
            _seed_data_fallback(conn)
            logger.info(f"[manufacturing_tools] Excel 不可用，使用回退模拟数据: {_DB_PATH}")

        conn.close()
        _db_ready = True
        return True
    except Exception as e:
        logger.error(f"[manufacturing_tools] 数据库初始化失败: {e}")
        return False


def _create_tables(conn: sqlite3.Connection):
    conn.executescript("""
        CREATE TABLE orders (
            order_id TEXT PRIMARY KEY,
            customer TEXT NOT NULL,
            product TEXT NOT NULL,
            quantity REAL NOT NULL,
            unit_price REAL,
            status TEXT NOT NULL,
            order_date TEXT NOT NULL,
            delivery_date TEXT,
            responsible TEXT
        );
        CREATE TABLE inventory (
            material_code TEXT PRIMARY KEY,
            material_name TEXT NOT NULL,
            spec TEXT,
            quantity REAL NOT NULL,
            unit TEXT NOT NULL,
            warehouse TEXT NOT NULL,
            last_inbound TEXT,
            safety_stock REAL DEFAULT 0,
            status TEXT DEFAULT '正常'
        );
        CREATE TABLE quality_inspections (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            batch_no TEXT NOT NULL,
            material_name TEXT NOT NULL,
            inspection_item TEXT NOT NULL,
            result TEXT NOT NULL,
            inspector TEXT NOT NULL,
            inspection_date TEXT NOT NULL,
            remark TEXT
        );
    """)


def _find_excel(filename: str) -> str:
    """在 data 目录查找 Excel 文件"""
    candidates = [_DATA_DIR / filename, Path("/app/data") / filename]
    for p in candidates:
        if p.exists():
            return str(p)
    return ""


def _load_from_excel(conn: sqlite3.Connection) -> bool:
    """从 Excel 文件加载数据到 SQLite。成功返回 True。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        logger.warning("[manufacturing_tools] openpyxl 未安装，无法读取 Excel")
        return False

    order_file = _find_excel("订单数据.xlsx")
    inventory_file = _find_excel("库存数据.xlsx")
    quality_file = _find_excel("质检数据.xlsx")

    if not all([order_file, inventory_file, quality_file]):
        missing = []
        if not order_file:
            missing.append("订单数据.xlsx")
        if not inventory_file:
            missing.append("库存数据.xlsx")
        if not quality_file:
            missing.append("质检数据.xlsx")
        logger.warning(f"[manufacturing_tools] Excel 文件缺失: {missing}")
        return False

    try:
        _load_orders_from_excel(conn, order_file)
        _load_inventory_from_excel(conn, inventory_file)
        _load_quality_from_excel(conn, quality_file)
        conn.commit()
        return True
    except Exception as e:
        conn.rollback()
        logger.error(f"[manufacturing_tools] Excel 数据加载失败: {e}")
        return False


def _read_excel_rows(filepath: str):
    """读取 Excel 文件，返回 (headers, rows)"""
    from openpyxl import load_workbook
    wb = load_workbook(filepath, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return [], []
    header = [str(h or "").strip() for h in rows[0]]
    return header, rows[1:]


def _col_index(header: list, keyword: str) -> int:
    """按关键词模糊匹配列索引"""
    for i, h in enumerate(header):
        if keyword in h:
            return i
    return -1


def _cell(row, idx, default=""):
    """安全取单元格值，datetime 自动格式化为 YYYY-MM-DD"""
    if idx < 0 or idx >= len(row):
        return default
    v = row[idx]
    if v is None:
        return default
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return v


def _load_orders_from_excel(conn: sqlite3.Connection, filepath: str):
    header, rows = _read_excel_rows(filepath)
    c_id = _col_index(header, "订单号")
    c_cust = _col_index(header, "客户名称")
    c_prod = _col_index(header, "产品名称")
    c_qty = _col_index(header, "数量")
    c_price = _col_index(header, "单价")
    c_status = _col_index(header, "状态")
    c_order = _col_index(header, "下单日期")
    c_delivery = _col_index(header, "交货日期")
    c_resp = _col_index(header, "负责人")

    count = 0
    for row in rows:
        order_id = str(_cell(row, c_id, "")).strip()
        if not order_id:
            continue
        conn.execute(
            "INSERT OR IGNORE INTO orders VALUES (?,?,?,?,?,?,?,?,?)",
            (
                order_id,
                str(_cell(row, c_cust)),
                str(_cell(row, c_prod)),
                float(_cell(row, c_qty, 0)),
                float(_cell(row, c_price, 0)),
                str(_cell(row, c_status)),
                str(_cell(row, c_order)),
                str(_cell(row, c_delivery)),
                str(_cell(row, c_resp)),
            ),
        )
        count += 1
    logger.info(f"[manufacturing_tools] 订单 Excel 加载 {count} 条 ({filepath})")


def _load_inventory_from_excel(conn: sqlite3.Connection, filepath: str):
    header, rows = _read_excel_rows(filepath)
    c_code = _col_index(header, "物料编码")
    c_name = _col_index(header, "物料名称")
    c_spec = _col_index(header, "规格型号")
    c_qty = _col_index(header, "库存数量")
    c_unit = _col_index(header, "单位")
    c_wh = _col_index(header, "仓库")
    c_safety = _col_index(header, "安全库存")
    c_inbound = _col_index(header, "最近入库日期")
    c_status = _col_index(header, "状态")

    count = 0
    for row in rows:
        code = str(_cell(row, c_code, "")).strip()
        if not code:
            continue
        conn.execute(
            "INSERT OR IGNORE INTO inventory VALUES (?,?,?,?,?,?,?,?,?)",
            (
                code,
                str(_cell(row, c_name)),
                str(_cell(row, c_spec)),
                float(_cell(row, c_qty, 0)),
                str(_cell(row, c_unit)),
                str(_cell(row, c_wh)),
                str(_cell(row, c_inbound)),
                float(_cell(row, c_safety, 0)),
                str(_cell(row, c_status)),
            ),
        )
        count += 1
    logger.info(f"[manufacturing_tools] 库存 Excel 加载 {count} 条 ({filepath})")


def _load_quality_from_excel(conn: sqlite3.Connection, filepath: str):
    header, rows = _read_excel_rows(filepath)
    c_batch = _col_index(header, "批次号")
    c_product = _col_index(header, "产品名称")
    c_item = _col_index(header, "检验项目")
    c_result = _col_index(header, "判定")
    c_inspector = _col_index(header, "检验员")
    c_date = _col_index(header, "检验日期")
    c_remark = _col_index(header, "备注")

    count = 0
    for row in rows:
        batch = str(_cell(row, c_batch, "")).strip()
        if not batch:
            continue
        conn.execute(
            "INSERT INTO quality_inspections (batch_no, material_name, inspection_item, result, inspector, inspection_date, remark) VALUES (?,?,?,?,?,?,?)",
            (
                batch,
                str(_cell(row, c_product)),
                str(_cell(row, c_item)),
                str(_cell(row, c_result, "合格")),
                str(_cell(row, c_inspector)),
                str(_cell(row, c_date)),
                str(_cell(row, c_remark, "")),
            ),
        )
        count += 1
    logger.info(f"[manufacturing_tools] 质检 Excel 加载 {count} 条 ({filepath})")


def _seed_data_fallback(conn: sqlite3.Connection):
    """Excel 不可用时的回退: 使用内存模拟数据。"""
    import random
    random.seed(42)
    base_date = datetime(2025, 1, 1)

    people = ["陈经理", "刘主管", "赵经理", "孙主管", "周经理"]
    for i in range(1, 51):
        order_date = base_date + timedelta(days=random.randint(0, 500))
        status = random.choice(_ORDER_STATUSES)
        delivery = order_date + timedelta(days=random.randint(15, 90))
        conn.execute(
            "INSERT INTO orders VALUES (?,?,?,?,?,?,?,?,?)",
            (
                f"WO{2025000 + i}",
                random.choice(_CUSTOMERS),
                random.choice(_PRODUCTS),
                round(random.uniform(1, 100), 1),
                round(random.uniform(5000, 50000), 2),
                status,
                order_date.strftime("%Y-%m-%d"),
                delivery.strftime("%Y-%m-%d"),
                random.choice(people),
            ),
        )

    for code, name, unit, base_qty in _MATERIALS:
        qty = round(base_qty * random.uniform(0.3, 1.5), 1)
        last_in = base_date + timedelta(days=random.randint(30, 500))
        conn.execute(
            "INSERT INTO inventory VALUES (?,?,?,?,?,?,?,?,?)",
            (
                code, name, "标准规格", qty, unit,
                random.choice(_WAREHOUSES),
                last_in.strftime("%Y-%m-%d"),
                round(base_qty * 0.3, 1),
                "正常",
            ),
        )

    for i in range(1, 101):
        product = random.choice(_PRODUCTS)
        item = random.choice(_INSPECTION_ITEMS.get(product, ["外观"]))
        result = random.choices(["合格", "不合格"], weights=[85, 15])[0]
        insp_date = base_date + timedelta(days=random.randint(0, 500))
        conn.execute(
            "INSERT INTO quality_inspections (batch_no, material_name, inspection_item, result, inspector, inspection_date, remark) VALUES (?,?,?,?,?,?,?)",
            (
                f"B{20250000 + i}",
                product, item, result,
                random.choice(_INSPECTORS),
                insp_date.strftime("%Y-%m-%d"),
                "无异常" if result == "合格" else random.choice(["指标偏高", "指标偏低", "外观异常"]),
            ),
        )

    conn.commit()


# ═══════════════════════════════════════════════════════════════
#  工具函数
# ═══════════════════════════════════════════════════════════════


def tool_query_order(order_id: str = "", status: str = "", customer: str = "") -> str:
    """查询订单信息。支持按订单号精确查询，或按状态/客户筛选。"""
    if not _ensure_db():
        return json.dumps({"error": "订单数据库不可用"}, ensure_ascii=False)
    try:
        conn = sqlite3.connect(_DB_PATH)
        conn.row_factory = sqlite3.Row
        sql = "SELECT * FROM orders WHERE 1=1"
        params = []
        if order_id.strip():
            sql += " AND order_id = ?"
            params.append(order_id.strip())
        if status.strip():
            sql += " AND status = ?"
            params.append(status.strip())
        if customer.strip():
            sql += " AND customer LIKE ?"
            params.append(f"%{customer.strip()}%")
        sql += " ORDER BY order_date DESC LIMIT 20"
        rows = conn.execute(sql, params).fetchall()
        conn.close()
        data = [dict(r) for r in rows]
        return json.dumps({"count": len(data), "orders": data}, ensure_ascii=False)
    except Exception as e:
        logger.error(f"[manufacturing_tools] 订单查询失败: {e}")
        return json.dumps({"error": f"查询失败: {e}"}, ensure_ascii=False)


def tool_query_inventory(material_code: str = "", warehouse: str = "", material_name: str = "") -> str:
    """查询物料库存信息。支持按物料编码、仓库名称或物料名称查询。"""
    if not _ensure_db():
        return json.dumps({"error": "库存数据库不可用"}, ensure_ascii=False)
    try:
        conn = sqlite3.connect(_DB_PATH)
        conn.row_factory = sqlite3.Row
        sql = "SELECT * FROM inventory WHERE 1=1"
        params = []
        if material_code.strip():
            sql += " AND material_code = ?"
            params.append(material_code.strip())
        if warehouse.strip():
            sql += " AND warehouse LIKE ?"
            params.append(f"%{warehouse.strip()}%")
        if material_name.strip():
            sql += " AND material_name LIKE ?"
            params.append(f"%{material_name.strip()}%")
        sql += " ORDER BY material_code"
        rows = conn.execute(sql, params).fetchall()
        conn.close()
        data = []
        for r in rows:
            d = dict(r)
            d["below_safety"] = d["quantity"] < d.get("safety_stock", 0)
            data.append(d)
        return json.dumps({"count": len(data), "inventory": data}, ensure_ascii=False)
    except Exception as e:
        logger.error(f"[manufacturing_tools] 库存查询失败: {e}")
        return json.dumps({"error": f"查询失败: {e}"}, ensure_ascii=False)


def tool_query_quality(batch_no: str = "", result_filter: str = "", material_name: str = "", days: int = 0) -> str:
    """查询质检记录。支持按批次号、检验结果（合格/不合格）、物料名称筛选，或查询最近N天的记录。"""
    if not _ensure_db():
        return json.dumps({"error": "质检数据库不可用"}, ensure_ascii=False)
    try:
        conn = sqlite3.connect(_DB_PATH)
        conn.row_factory = sqlite3.Row
        sql = "SELECT * FROM quality_inspections WHERE 1=1"
        params = []
        if batch_no.strip():
            sql += " AND batch_no = ?"
            params.append(batch_no.strip())
        if result_filter.strip():
            sql += " AND result = ?"
            params.append(result_filter.strip())
        if material_name.strip():
            sql += " AND material_name LIKE ?"
            params.append(f"%{material_name.strip()}%")
        if days > 0:
            cutoff = (datetime.now() - timedelta(days=days)).strftime("%Y-%m-%d")
            sql += " AND inspection_date >= ?"
            params.append(cutoff)
        sql += " ORDER BY inspection_date DESC LIMIT 30"
        rows = conn.execute(sql, params).fetchall()
        conn.close()
        data = [dict(r) for r in rows]
        return json.dumps({"count": len(data), "inspections": data}, ensure_ascii=False)
    except Exception as e:
        logger.error(f"[manufacturing_tools] 质检查询失败: {e}")
        return json.dumps({"error": f"查询失败: {e}"}, ensure_ascii=False)